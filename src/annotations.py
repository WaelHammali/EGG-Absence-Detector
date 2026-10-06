"""Resolve provisional or official annotation sources into one normalized schema."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json
import math

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import yaml

from src.io import FS, SCORING_CHANNELS, load_recording


@dataclass
class WorkbookRow:
    recording: str
    source_row: int
    trace_clock: float
    intervals: list[tuple[float, float, int]]
    intercritical: list[tuple[float, int]]


def _clock_seconds(value) -> float:
    if value is None:
        raise ValueError("Missing clock value")
    if isinstance(value, str):
        parts = [float(part) for part in value.split(":")]
        if len(parts) != 3:
            raise ValueError(f"Invalid clock string: {value}")
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    return float(value) * 86400


def _relative(clock: float, trace_clock: float) -> float:
    """Clock difference with midnight rollover."""
    value = clock - trace_clock
    return value + 86400 if value < 0 else value


def load_decisions(root: Path, path: str | Path | None = None) -> dict:
    path = Path(path) if path else root / "config/annotation_decisions.yaml"
    with path.open() as handle:
        return yaml.safe_load(handle)


def read_official_workbook(path: str | Path) -> dict[str, list[WorkbookRow]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook.active
    output: dict[str, list[WorkbookRow]] = {}
    populated_rows = sorted({row for row, _ in sheet._cells})
    for row_number in populated_rows:
        raw_id = sheet.cell(row_number, 1).value
        if row_number == 1 or raw_id is None:
            continue
        recording = str(raw_id).strip().replace("-", "_")
        trace_value = sheet.cell(row_number, 6).value
        if trace_value is None:
            raise ValueError(f"{recording} row {row_number}: DebutTrace is missing")
        trace_clock = _clock_seconds(trace_value)
        intervals = []
        for number in range(1, 13):
            start_value = sheet.cell(row_number, 7 + (number - 1) * 2).value
            end_value = sheet.cell(row_number, 8 + (number - 1) * 2).value
            if start_value is None and end_value is None:
                continue
            if start_value is None or end_value is None:
                raise ValueError(f"{recording} row {row_number}: incomplete CE{number}")
            intervals.append((_clock_seconds(start_value), _clock_seconds(end_value), number))
        intercritical = []
        for number in range(1, 8):
            value = sheet.cell(row_number, 31 + number - 1).value
            if value is not None:
                intercritical.append((_clock_seconds(value), number))
        output.setdefault(recording, []).append(
            WorkbookRow(recording, row_number, trace_clock, intervals, intercritical)
        )
    workbook.close()
    return output


def _relative_power(signal: np.ndarray, scoring: dict) -> tuple[np.ndarray, np.ndarray]:
    window_samples = round(float(scoring["window_seconds"]) * FS)
    hop_samples = round(float(scoring["hop_seconds"]) * FS)
    count = 1 + (len(signal) - window_samples) // hop_samples
    hann = np.hanning(window_samples).astype(np.float32)
    frequencies = np.fft.rfftfreq(window_samples, d=1 / FS)
    low, high = map(float, scoring["target_band_hz"])
    total_low, total_high = map(float, scoring["denominator_band_hz"])
    target = (frequencies >= low) & (frequencies <= high)
    total = (frequencies >= total_low) & (frequencies <= total_high)
    result = np.empty(count, dtype=np.float32)
    samples = np.arange(window_samples)
    for first in range(0, count, 512):
        last = min(first + 512, count)
        starts = np.arange(first, last) * hop_samples
        windows = signal[starts[:, None] + samples[None, :], :]
        windows = windows - windows.mean(axis=1, keepdims=True)
        spectrum = np.fft.rfft(windows * hann[None, :, None], axis=1)
        power = spectrum.real * spectrum.real + spectrum.imag * spectrum.imag
        ratios = power[:, target, :].sum(axis=1) / np.maximum(power[:, total, :].sum(axis=1), 1e-20)
        result[first:last] = ratios.mean(axis=1)
    centers = float(scoring["window_seconds"]) / 2 + np.arange(count) * float(scoring["hop_seconds"])
    return centers, result


def score_intervals(
    signal: np.ndarray,
    intervals: list[tuple[float, float]],
    scoring: dict,
) -> dict:
    times, values = _relative_power(signal, scoring)
    mask = np.zeros(len(times), dtype=bool)
    for start, end in intervals:
        mask |= (times >= start) & (times <= end)
    if not mask.any() or mask.all():
        return {"score": -math.inf, "inside_outside_ratio": math.nan, "inside_windows": int(mask.sum())}
    inside = float(values[mask].mean())
    outside = float(values[~mask].mean())
    return {
        "score": inside - outside,
        "inside_outside_ratio": inside / max(outside, 1e-20),
        "inside_windows": int(mask.sum()),
    }


def _row_intervals(row: WorkbookRow) -> list[tuple[float, float, int]]:
    return [(_relative(start, row.trace_clock), _relative(end, row.trace_clock), number) for start, end, number in row.intervals]


def resolve_annotations(
    root: Path,
    mode: str,
    output_dir: Path,
    decisions_path: str | Path | None = None,
) -> dict:
    """Resolve a source and write normalized annotations, offsets and exclusions."""
    decisions = load_decisions(root, decisions_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    if mode == "provisional":
        return _resolve_provisional(root, output_dir, decisions)
    if mode != "official":
        raise ValueError("mode must be 'official' or 'provisional'")
    return _resolve_official(root, output_dir, decisions)


def _resolve_provisional(root: Path, output_dir: Path, decisions: dict) -> dict:
    sources = decisions["sources"]
    annotations = pd.read_csv(root / sources["provisional_annotations"])
    offsets = pd.read_csv(root / sources["provisional_offsets"])
    intercritical = pd.read_csv(root / sources["provisional_intercritical"])
    excluded = decisions["provisional"]["excluded_recordings"]
    pd.DataFrame(
        [{"recording": key, "reason": value} for key, value in excluded.items()]
    ).to_csv(output_dir / "exclusions.csv", index=False)
    annotations.to_csv(output_dir / "annotations_clean.csv", index=False)
    offsets.to_csv(output_dir / "offsets.csv", index=False)
    intercritical.to_csv(output_dir / "intercritical_events.csv", index=False)
    return {
        "annotations": annotations,
        "offsets": offsets,
        "intercritical": intercritical,
        "excluded": excluded,
        "decisions": ["Provisional inferred annotations copied without modification."],
        "scores": {},
    }


def _resolve_official(root: Path, output_dir: Path, decisions: dict) -> dict:
    config = decisions["official"]
    scoring = config["scoring"]
    sources = decisions["sources"]
    workbook_rows = read_official_workbook(root / sources["official_workbook"])
    provisional_annotations = pd.read_csv(root / sources["provisional_annotations"])
    provisional_offsets = pd.read_csv(root / sources["provisional_offsets"])
    provisional_intercritical = pd.read_csv(root / sources["provisional_intercritical"])
    raw_summary = json.loads((root / "docs/inspection_summary.json").read_text())
    durations = {row["Recording"]: float(row["duration_s"]) for row in raw_summary["recordings"]}

    excluded = dict(config["excluded_recordings"])
    selected: dict[str, WorkbookRow] = {}
    score_log: dict[str, list[dict]] = {}
    decision_log: list[str] = []

    duplicate_config = config.get("duplicate_recordings", {})
    for recording, rows in workbook_rows.items():
        if recording in excluded:
            continue
        if len(rows) == 1:
            selected[recording] = rows[0]
            continue
        rule = duplicate_config.get(recording)
        if rule is None:
            raise ValueError(f"{recording}: duplicate workbook rows need a YAML decision")
        allowed = set(rule.get("candidate_rows", [row.source_row for row in rows]))
        candidates = [row for row in rows if row.source_row in allowed]
        strategy = rule["strategy"]
        if strategy == "best_signal_score":
            frame = load_recording(root / f"data/raw/{recording}_0000d.mat")
            signal = frame.loc[:, list(SCORING_CHANNELS)].to_numpy(dtype=np.float32, copy=False)
            scored = []
            for row in candidates:
                intervals = [(start, end) for start, end, _ in _row_intervals(row)]
                metrics = score_intervals(signal, intervals, scoring)
                metrics.update({"source_row": row.source_row})
                scored.append(metrics)
            chosen_metrics = max(scored, key=lambda value: value["score"])
            selected[recording] = next(row for row in candidates if row.source_row == chosen_metrics["source_row"])
            score_log[recording] = scored
            decision_log.append(
                f"{recording}: duplicate rows scored; selected row {chosen_metrics['source_row']} "
                f"(score={chosen_metrics['score']:.6f}, ratio={chosen_metrics['inside_outside_ratio']:.3f})."
            )
        elif strategy == "row":
            chosen_row = int(rule["row"])
            selected[recording] = next(row for row in candidates if row.source_row == chosen_row)
            decision_log.append(f"{recording}: duplicate row {chosen_row} selected explicitly by YAML.")
        else:
            raise ValueError(f"{recording}: unsupported duplicate strategy {strategy}")

    annotation_rows = []
    offset_rows = []
    intercritical_rows = []
    correction_config = config.get("corrections", {})
    for recording, row in selected.items():
        intervals = _row_intervals(row)
        correction = correction_config.get(recording)
        if correction:
            number = int(correction["seizure_number"])
            replacement_end = _clock_seconds(correction["replacement_end_clock"])
            replacement_relative = _relative(replacement_end, row.trace_clock)
            intervals = [
                (start, replacement_relative if source_number == number else end, source_number)
                for start, end, source_number in intervals
            ]
            candidate = [(start, end) for start, end, source_number in intervals if source_number == number]
            strategy = correction["strategy"]
            keep = strategy == "always_correct"
            metrics = None
            if strategy == "signal_confirmed":
                frame = load_recording(root / f"data/raw/{recording}_0000d.mat")
                signal = frame.loc[:, list(SCORING_CHANNELS)].to_numpy(dtype=np.float32, copy=False)
                metrics = score_intervals(signal, candidate, scoring)
                keep = metrics["score"] > 0 and metrics["inside_outside_ratio"] >= float(correction["minimum_inside_outside_ratio"])
                score_log[f"{recording}_correction"] = [metrics]
            elif strategy == "drop":
                keep = False
            if not keep:
                intervals = [value for value in intervals if value[2] != number]
            decision_log.append(
                f"{recording} CE{number}: proposed end {correction['replacement_end_clock']} "
                + (
                    f"scored {metrics['score']:.6f}, ratio={metrics['inside_outside_ratio']:.3f}; "
                    if metrics else ""
                )
                + ("correction retained." if keep else "interval dropped.")
            )

        duration = durations.get(recording)
        for seizure_id, (start, end, source_number) in enumerate(intervals, 1):
            if end <= start:
                raise ValueError(f"{recording} row {row.source_row} CE{source_number}: end <= start")
            if duration is not None and (start < 0 or end > duration):
                raise ValueError(
                    f"{recording} row {row.source_row} CE{source_number}: interval {start}–{end} outside duration {duration}"
                )
            annotation_rows.append(
                {
                    "recording": recording,
                    "seizure_id": seizure_id,
                    "start_s": round(start, 6),
                    "end_s": round(end, 6),
                    "duration_s": round(end - start, 6),
                    "source_row": row.source_row,
                    "note": "Official DebutTrace alignment from Annotations MAJ.xlsx.",
                }
            )
        offset_rows.append(
            {
                "recording": recording,
                "offset_clock": round(row.trace_clock, 6),
                "start_clock": _format_clock(row.trace_clock),
                "method": "official",
                "confidence": "NA",
                "agrees_with_header": "NA",
            }
        )
        for event_id, (clock, source_number) in enumerate(row.intercritical, 1):
            elapsed = _relative(clock, row.trace_clock)
            intercritical_rows.append(
                {
                    "recording": recording,
                    "event_id": event_id,
                    "clock_s": round(clock, 6),
                    "event_s": round(elapsed, 6),
                    "Class": 0,
                    "source_row": row.source_row,
                    "source_column": f"Intercrit{source_number}",
                    "inside_recording": "yes" if 0 <= elapsed <= durations.get(recording, math.inf) else "no",
                    "note": "Official clock point; interictal, not a seizure.",
                }
            )

    for recording, rule in config.get("missing_recordings", {}).items():
        if rule["strategy"] == "exclude":
            excluded[recording] = rule.get("note", "Missing from official workbook.")
            decision_log.append(f"{recording}: missing from official workbook and excluded by YAML.")
            continue
        if rule["strategy"] != "fallback_provisional":
            raise ValueError(f"{recording}: unsupported missing strategy {rule['strategy']}")
        fallback = provisional_annotations.loc[provisional_annotations["recording"] == recording].copy()
        if fallback.empty:
            raise ValueError(f"{recording}: no provisional intervals available for fallback")
        fallback["note"] = "Fallback: " + rule["note"]
        annotation_rows.extend(fallback.to_dict("records"))
        offset = provisional_offsets.loc[provisional_offsets["recording"] == recording].iloc[0].to_dict()
        offset["method"] = "fallback"
        offset_rows.append(offset)
        fallback_events = provisional_intercritical.loc[provisional_intercritical["recording"] == recording]
        intercritical_rows.extend(fallback_events.to_dict("records"))
        decision_log.append(f"{recording}: missing from official workbook; provisional intervals retained as fallback.")

    for recording in config.get("validation_recordings", []):
        row = selected.get(recording)
        if row is None:
            continue
        frame = load_recording(root / f"data/raw/{recording}_0000d.mat")
        signal = frame.loc[:, list(SCORING_CHANNELS)].to_numpy(dtype=np.float32, copy=False)
        intervals = [(start, end) for start, end, _ in _row_intervals(row)]
        metrics = score_intervals(signal, intervals, scoring)
        score_log[f"{recording}_official_validation"] = [metrics]
        decision_log.append(
            f"{recording}: official DebutTrace places the interval at "
            f"{intervals[0][0]:.3f}–{intervals[0][1]:.3f} s; "
            f"score={metrics['score']:.6f}, ratio={metrics['inside_outside_ratio']:.3f}; included."
        )

    annotations = pd.DataFrame(annotation_rows).sort_values(["recording", "seizure_id"]).reset_index(drop=True)
    offsets = pd.DataFrame(offset_rows).sort_values("recording").reset_index(drop=True)
    intercritical = pd.DataFrame(intercritical_rows)
    exclusion_frame = pd.DataFrame([{"recording": key, "reason": value} for key, value in excluded.items()])
    annotations.to_csv(output_dir / "annotations_clean.csv", index=False)
    offsets.to_csv(output_dir / "offsets.csv", index=False)
    intercritical.to_csv(output_dir / "intercritical_events.csv", index=False)
    exclusion_frame.to_csv(output_dir / "exclusions.csv", index=False)
    return {
        "annotations": annotations,
        "offsets": offsets,
        "intercritical": intercritical,
        "excluded": excluded,
        "decisions": decision_log,
        "scores": score_log,
    }


def _format_clock(seconds: float) -> str:
    seconds %= 86400
    return f"{int(seconds // 3600):02d}:{int(seconds % 3600 // 60):02d}:{int(round(seconds % 60)):02d}"
