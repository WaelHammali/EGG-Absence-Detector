"""Read-only audit of raw EEG files; write compact evidence outside data/raw."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import scipy.io
import h5py
from openpyxl import load_workbook
from matio import load_from_mat

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / 'data/raw'


def load_timetable(path, variable):
    """Try SciPy, HDF5, then mat-io for MATLAB MCOS timetables."""
    try:
        basic = scipy.io.loadmat(path, squeeze_me=True, struct_as_record=False)
        variables = [k for k in basic if not k.startswith('__')]
        if isinstance(basic.get(variable), pd.DataFrame):
            return basic[variable], 'scipy', variables
        del basic
    except NotImplementedError:
        with h5py.File(path, 'r') as handle:
            variables = list(handle.keys())
    table = load_from_mat(path, add_table_attrs=True)[variable]
    if not isinstance(table, pd.DataFrame):
        raise TypeError(f'{path.name}: expected a decoded timetable')
    return table, 'mat-io', variables


def annotations():
    # Normal mode reads actual XML cells without iterating a million styled rows.
    wb = load_workbook(DATA_DIR / 'Annotations.xlsx', data_only=True)
    result = {}
    for sheet in wb:
        populated = sorted({r for (r, c), cell in sheet._cells.items()
                            if cell.value is not None})
        headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
        rows = []
        totals = []
        for r in populated:
            if r == 1:
                continue
            values = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            if values[0] is None:
                totals.append({'excel_row': r, 'values': values})
                continue
            intervals = []
            for i in range(12):
                start, end = values[5 + i * 2:7 + i * 2]
                if start is not None or end is not None:
                    intervals.append({'number': i + 1, 'start_day': start, 'end_day': end,
                                      'duration_s': None if start is None or end is None
                                      else round((end - start) * 86400, 6)})
            rows.append({'excel_row': r, 'Recording': str(values[0]).strip().replace('-', '_'),
                         'Sexe': values[1], 'Age': values[2], 'NombreCrises': values[3],
                         'TypeAbs': values[4], 'intervals': intervals,
                         'intercritical': values[29:]})
        result[sheet.title] = {'declared_dimensions': [sheet.max_row, sheet.max_column],
                               'headers': headers, 'rows': rows, 'totals': totals}
    wb.close()
    return result


def main():
    report = {'recordings': [], 'annotations': annotations()}
    d_files = sorted(DATA_DIR.glob('*d.mat'))
    h_files = sorted(DATA_DIR.glob('*h.mat'))
    if len(d_files) != len(h_files):
        raise ValueError('Unmatched file counts')
    for path in d_files:
        recording = path.name.removesuffix('_0000d.mat')
        table, loader, variables = load_timetable(path, 'output_d')
        starts = table.index.total_seconds().to_numpy()
        if not np.allclose(np.diff(starts), 1):
            raise ValueError(f'{recording}: irregular record timestamps')
        stats = {}
        for channel in table.columns:
            lengths = {np.asarray(cell).size for cell in table[channel]}
            if len(lengths) != 1:
                raise ValueError(f'{recording}: uneven blocks in {channel}')
            signal = np.concatenate([np.asarray(cell).reshape(-1) for cell in table[channel]])
            finite = signal[np.isfinite(signal)]
            stats[channel] = {'block_samples': next(iter(lengths)), 'samples': len(signal),
                              'dtype': str(signal.dtype), 'nan': int(np.isnan(signal).sum()),
                              'infinite': int(np.isinf(signal).sum()),
                              'min': float(finite.min()) if finite.size else None,
                              'max': float(finite.max()) if finite.size else None,
                              'constant': bool(finite.size and finite.min() == finite.max()),
                              'first_samples': signal[:5].tolist()}
        block_sizes = {x['block_samples'] for x in stats.values()}
        if len(block_sizes) != 1:
            raise ValueError(f'{recording}: differing sampling frequencies')
        fs = next(iter(block_sizes)) / float(np.diff(starts)[0])
        hdr, _, h_variables = load_timetable(path.with_name(path.name.replace('d.mat','h.mat')), 'output_h')
        events = [{'onset_s': float(t.total_seconds()), 'text': str(row['Annotations']),
                   'duration_s': None if pd.isna(row['Duration']) else float(row['Duration'].total_seconds())}
                  for t, row in hdr.iterrows()]
        item = {'Recording': recording, 'loader': loader, 'd_variables': variables,
                'h_variables': h_variables, 'timetable_shape': list(table.shape),
                'index_name': table.index.name, 'start_s': float(starts[0]),
                'duration_s': float(starts[-1] + 1 - starts[0]), 'fs_hz': fs,
                'samples': next(iter(stats.values()))['samples'], 'channels': stats,
                'units': table.attrs.get('varUnits', []),
                'header_shape': list(hdr.shape), 'header_events': events,
                'header_duration_missing': int(hdr['Duration'].isna().sum())}
        report['recordings'].append(item)
        print(f'{recording}: {item["samples"]:,} samples, {len(stats)} channels, {fs:g} Hz, '
              f'{item["duration_s"]:g} s; constants: {[c for c,s in stats.items() if s["constant"]]}', flush=True)
        del table, hdr, signal, finite
    out = ROOT / 'docs/inspection_summary.json'
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    rows = report['annotations']['Feuil1']['rows']
    print(f'Annotations: {len(rows)} rows, {len({r["Recording"] for r in rows})} unique recordings, '
          f'{sum(len(r["intervals"]) for r in rows)} populated interval pairs')
    print('Evidence:', out)


if __name__ == '__main__':
    main()
