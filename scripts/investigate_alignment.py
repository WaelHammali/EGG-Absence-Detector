"""Audit alignment without altering raw files or creating a learning dataset."""
from pathlib import Path
import base64
import json
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import scipy.io
from scipy.signal import spectrogram
from openpyxl import load_workbook
from matio import load_from_mat

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw'
DOCS = ROOT / 'docs'
FIGS = ROOT / 'outputs/figures/alignment'
SELECTED = ['191113A_D', '210204B_C', '210208B_G']


def encode(obj):
    """Include every raw MCOS field and array element, with shape/type metadata."""
    if hasattr(obj, 'properties'):
        return {'class': obj.classname, 'type_system': str(obj.type_system),
                'class_alias': obj.class_alias, 'properties': encode(obj.properties)}
    if isinstance(obj, np.ndarray):
        values = ([{name: encode(item[name]) for name in obj.dtype.names} for item in obj.flat]
                  if obj.dtype.names else [encode(item) for item in obj.flat])
        return {'dtype': str(obj.dtype), 'shape': list(obj.shape), 'values': values}
    if isinstance(obj, dict):
        return {key: encode(value) for key, value in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [encode(value) for value in obj]
    if isinstance(obj, bytes):
        return {'base64': base64.b64encode(obj).decode(), 'text': obj.decode(errors='replace')}
    if isinstance(obj, np.generic):
        return encode(obj.item())
    if isinstance(obj, float) and not np.isfinite(obj):
        return repr(obj)
    return obj


def main():
    FIGS.mkdir(parents=True, exist_ok=True)
    summary = json.loads((DOCS / 'inspection_summary.json').read_text())
    recordings = {r['Recording']: r for r in summary['recordings']}
    rows = summary['annotations']['Feuil1']['rows']
    wb = load_workbook(RAW / 'Annotations.xlsx', data_only=False)
    sheet = wb['Feuil1']
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(RAW / 'Annotations.xlsx') as z:
        xml = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    xml_cells = {c.attrib['r']: c for c in xml.findall('.//s:sheetData/s:row/s:c', ns)}
    lines = ['# Time alignment investigation', '',
             'Raw inputs are unchanged. No labels, conversions, dataset or model have been built.', '',
             '## 1. Exact workbook cells and complete header', '',
             'Selected records: `191113A_D` (Excel row 3), `210204B_C` (row 5), '
             '`210208B_G` (row 6), each with at least two populated intervals. '
             'The first three ID rows of the workbook are `190304A-E`, `191113A-D`, '
             '` 200625A-F `; the third has no seizure pairs.', '',
             'Values below retain the original XLSX XML numeric text; the Python value is shown '
             'with `repr`. Excel type is openpyxl `data_type` (`n`: numeric or empty, '
             '`s`: string); XML type is shown separately. An absent XML cell is empty, '
             'not a recorded zero. All selected cells use number format `General`.', '']
    cells_report = []
    # Include all 36 headers and all cells of each selected row, even blanks.
    for row_number in [1, 3, 5, 6]:
        lines += [f'### Excel row {row_number}', '',
                  '| Cell | Column header | Python value | Excel type | XML type | Exact XML value/text |',
                  '|---|---|---|---|---|---|']
        for cell in sheet[row_number]:
            node = xml_cells.get(cell.coordinate)
            v = node.find('s:v', ns) if node is not None else None
            inline = node.find('s:is', ns) if node is not None else None
            literal = (v.text if v is not None else ''.join(inline.itertext()) if inline is not None else None)
            xtype = node.attrib.get('t', 'n (implicit)') if node is not None else 'absent'
            header = sheet.cell(1, cell.column).value
            record = {'cell': cell.coordinate, 'header': header, 'value': cell.value,
                      'python_repr': repr(cell.value), 'excel_type': cell.data_type,
                      'xml_type': xtype, 'xml_literal': literal, 'number_format': cell.number_format}
            cells_report.append(record)
            lines.append(f'| {cell.coordinate} | {header!r} | `{cell.value!r}` | {cell.data_type} | {xtype} | `{literal!r}` |')
    wb.close()
    (DOCS / 'alignment_raw_cells.json').write_text(json.dumps(cells_report, ensure_ascii=False, indent=2) + '\n')
    header_path = RAW / '190304A_E_0000h.mat'
    raw = load_from_mat(header_path, raw_data=True)
    basic = scipy.io.loadmat(header_path, squeeze_me=True, struct_as_record=False)
    hdr = load_from_mat(header_path, add_table_attrs=True)['output_h']
    header_dump = {'file': header_path.name,
                   'container_header': basic['__header__'].decode(),
                   'container_version': basic['__version__'], 'container_globals': basic['__globals__'],
                   'scipy_opaque': encode(np.asarray(basic['output_h'])),
                   'function_workspace': {'dtype': str(basic['__function_workspace__'].dtype),
                                          'shape': list(basic['__function_workspace__'].shape),
                                          'base64': base64.b64encode(basic['__function_workspace__'].tobytes()).decode()},
                   'raw_mcos': encode(raw), 'decoded_attrs': encode(hdr.attrs),
                   'decoded_index_name': hdr.index.name,
                   'decoded_events': recordings['190304A_E']['header_events']}
    (DOCS / 'alignment_header_complete.json').write_text(json.dumps(header_dump, ensure_ascii=False, indent=2) + '\n')
    fields = raw['output_h'].properties['any'].dtype.names
    lines += ['', '### Complete header: `190304A_E_0000h.mat`', '',
              'Every MCOS field and array value, all 77 decoded events, table attributes, '
              'container metadata and the complete binary internal workspace (base64) are in '
              '[alignment_header_complete.json](alignment_header_complete.json). No arrays are truncated.', '',
              f'Container description: `{header_dump["container_header"]}`. '
              'The creation time is a MAT-file export time, not the recording acquisition start.', '',
              'Raw timetable fields: ' + ', '.join(f'`{f}`' for f in fields) + '.', '',
              '`arrayProps` contains `Description` (empty), `UserData` (empty), '
              '`TableCustomProperties` (empty). `CustomProps` and `VariableCustomProps` are empty. '
              '`rowTimes` is a MATLAB `duration` object with relative elapsed seconds, not a `datetime`. '
              'The nested `Duration` data are missing (NaT); the other data column is `Annotations`. '
              'There is no acquisition date, absolute start time, or time-zone field in this header.', '',
              '| Onset (s) | Annotations (decoded, exact text) | Duration |', '|---:|---|---|']
    for e in recordings['190304A_E']['header_events']:
        lines.append(f'| {e["onset_s"]:g} | {e["text"]} | {e["duration_s"]} |')
    lines += ['', '## 2. Elapsed-time hypothesis, all recordings', '',
              'Test: interpret each numeric annotation as a duration in days, so '
              '`elapsed_seconds = Excel_value × 86400`, then require every interval to lie '
              'in `[0, recording_duration]`. No offset is subtracted. '
              'Max time includes every populated seizure endpoint and both versions of a duplicate; '
              '`Intercrit*` values are not seizure endpoints.', '',
              '| Recording | Duration (s) | Maximum annotated elapsed time (s) | Difference max − duration (s) | Consistent? |',
              '|---|---:|---:|---:|---|']
    tests = []
    for rid, rec in recordings.items():
        endpoints = [v[key] * 86400 for row in rows if row['Recording'] == rid
                     for v in row['intervals'] for key in ['start_day', 'end_day']]
        maximum = max(endpoints) if endpoints else None
        status = 'no annotations' if maximum is None else 'yes' if min(endpoints) >= 0 and maximum <= rec['duration_s'] else 'no'
        tests.append({'Recording': rid, 'duration_s': rec['duration_s'], 'max_annotation_s': maximum, 'status': status})
        lines.append(f'| {rid} | {rec["duration_s"]:g} | {maximum:.6f} | {maximum-rec["duration_s"]:.6f} | {status} |' if maximum is not None else f'| {rid} | {rec["duration_s"]:g} | — | — | no annotations |')
    yes = sum(t['status'] == 'yes' for t in tests)
    no = sum(t['status'] == 'no' for t in tests)
    lines += ['', f'**{yes}/22 recordings are positively consistent; {no}/22 are inconsistent; '
              '1/22 (`200625A_F`) cannot be tested because it has no annotations.** '
              'All 21 annotated recordings have every populated seizure interval starting '
              'after the recording has already ended. Thus even the first annotated seizures '
              'cannot be displayed at their proposed elapsed coordinates.', '',
              '## 3. Visual/Delta checks under that hypothesis', '',
              'Three figures include the real EEG Delta-power timeline and two 20 s views centered '
              'on the first two annotated seizure midpoints. Orange shading denotes proposed '
              'annotations, not validated seizures. The 20 s views contain no signal because '
              'their coordinates are outside the acquisition; this is explicitly marked rather '
              'than clipping timestamps or fabricating samples.', '',
              'Delta power is computed on all EEG channels and averaged across channels: '
              '2 s Hann windows, 1 s hop, mean removal per window, one-sided PSD density, '
              'sum of bins satisfying `0.5 ≤ f < 4 Hz` times the frequency-bin spacing '
              '(0.5 Hz). Units are signal amplitude², since physical units are unavailable. '
              'At this resolution 3 Hz is a Fourier bin. Supplementary first-20-s EEG/Delta '
              'plots show actual data; they have no seizure overlays because no proposed '
              'interval intersects this segment.', '',
              '**There are no samples or Delta windows inside any proposed interval. '
              'Inside/outside Delta ratios and seizure morphology at those times are undefined. '
              'The elapsed interpretation is rejected by coverage, not by interpreting an empty '
              'panel as absence of spike-and-wave activity.**', '']
    for rid in SELECTED:
        rec = recordings[rid]
        table = load_from_mat(RAW / f'{rid}_0000d.mat')['output_d']
        channels = [c for c in table if c.startswith('EEG')]
        signal = np.column_stack([np.concatenate([np.asarray(v).reshape(-1) for v in table[c]]) for c in channels]).astype('float32')
        fs = rec['fs_hz']
        f, times, psd = spectrogram(signal, fs=fs, window='hann', nperseg=int(2*fs),
                                    noverlap=int(fs), detrend='constant', scaling='density', axis=0)
        # PSD shape: frequency, channel, time.
        delta = psd[(f >= .5) & (f < 4)].sum(axis=0).mean(axis=0) * (f[1]-f[0])
        selected_intervals = next(r['intervals'] for r in rows if r['Recording'] == rid)[:2]
        fig, axes = plt.subplots(5, 1, figsize=(13, 13), constrained_layout=True)
        ax = axes[0]
        ax.plot(times, delta, color='navy', lw=.8, label='Observed EEG Delta power')
        ax.axvspan(0, rec['duration_s'], color='navy', alpha=.08, label='Recording coverage')
        for i, interval in enumerate(selected_intervals, 1):
            s, e = interval['start_day']*86400, interval['end_day']*86400
            ax.axvspan(s, e, color='orange', alpha=.7, label='Proposed annotation' if i == 1 else None)
        ax.set_xlim(0, max(v['end_day']*86400 for v in selected_intervals)+30)
        ax.set(title=f'{rid}: elapsed-time hypothesis versus observed Delta power',
               xlabel='Proposed elapsed time (s)', ylabel='Delta power (amplitude²)')
        ax.legend(loc='upper center')
        for i, interval in enumerate(selected_intervals):
            s, e = interval['start_day']*86400, interval['end_day']*86400
            midpoint = (s+e)/2
            for offset in [0,1]:
                ax=axes[1+2*i+offset]
                ax.axvspan(s,e,color='orange',alpha=.3)
                ax.set_xlim(midpoint-10,midpoint+10)
                ax.set(title=f'Proposed seizure {i+1}: {s:.3f}–{e:.3f} s; 20 s centered view',
                       xlabel='Proposed elapsed time (s)',
                       ylabel='EEGFp1 amplitude (unit unknown)' if offset==0 else 'Delta power (amplitude²)')
                ax.text(.5,.5,f'No samples: recording ends at {rec["duration_s"]:g} s',
                        ha='center',va='center',transform=ax.transAxes)
                ax.set_yticks([])
        fig.savefig(FIGS/f'{rid}_elapsed_alignment.png',dpi=160)
        plt.close(fig)
        fig,axes=plt.subplots(2,1,figsize=(13,6),constrained_layout=True)
        axes[0].plot(np.arange(int(20*fs))/fs,signal[:int(20*fs),channels.index('EEGFp1')],lw=.7)
        axes[0].set(title=f'{rid}: actual first 20 s, EEGFp1',xlabel='Elapsed time (s)',ylabel='Amplitude (unit unknown)')
        axes[1].plot(times[times<=20],delta[times<=20])
        axes[1].set(title='Actual Delta power averaged over EEG channels (no proposed seizure in this range)',
                    xlabel='Elapsed time (s)',ylabel='Delta power (amplitude²)')
        for ax in axes: ax.set_xlim(0,20)
        fig.savefig(FIGS/f'{rid}_actual_first20s.png',dpi=160)
        plt.close(fig)
        lines += [f'- `{rid}`: proposed intervals ' + ', '.join(f'{v["start_day"]*86400:.0f}–{v["end_day"]*86400:.0f} s' for v in selected_intervals)
                  + f'; recording ends at {rec["duration_s"]:g} s. '
                  + f'[Alignment plot](../outputs/figures/alignment/{rid}_elapsed_alignment.png), '
                  + f'[actual first 20 s](../outputs/figures/alignment/{rid}_actual_first20s.png).']
        del table, signal, psd
    lines += ['', '## 4. Annotation problems, without corrections', '',
              '### Exact explanation of 95 pairs versus declared 90', '',
              '`210914B_A`, Excel row 10, declares 6 but contains **11** pairs. '
              'The **five additional pairs are CE7–CE11** (ordinal positions beyond the declared six); '
              'this does not establish that those five are invalid.', '',
              '| Pair | Start cell and exact value | End cell and exact value | Time interpreted as hh:mm:ss | Duration (s) |',
              '|---|---|---|---|---:|']
    target=next(r for r in rows if r['Recording']=='210914B_A')
    def clock(days):
        s=int(round(days*86400));return f'{s//3600:02d}:{s%3600//60:02d}:{s%60:02d}'
    for v in target['intervals'][6:]:
        col=5+2*(v['number']-1)+1
        from openpyxl.utils import get_column_letter
        sc=f'{get_column_letter(col)}10'; ec=f'{get_column_letter(col+1)}10'
        lines.append(f'| CE{v["number"]} | {sc}: `{xml_cells[sc].find("s:v",ns).text}` | {ec}: `{xml_cells[ec].find("s:v",ns).text}` | {clock(v["start_day"])} → {clock(v["end_day"])} | {v["duration_s"]:g} |')
    lines += ['', '### Duplicate versions: no rows removed', '',
              'The earlier “93 after removing the duplicate” was a **hypothetical count**, '
              'not a deletion. No duplicates have been removed from any input. '
              '`230515B_G` occurs in Excel rows 15 and 22; both declare 2 seizures. '
              'Keeping one row would remove 2 pairs (95 → 93), and reduce the sum of declared '
              'counts from 90 to 88. These are conflicting versions, not exact duplicates.', '',
              '| Excel row | Age | Pair | Start raw value | End raw value | Time | Duration (s) |',
              '|---:|---:|---|---|---|---|---:|']
    for row in rows:
        if row['Recording']=='230515B_G':
            for v in row['intervals']:
                col=6+2*(v['number']-1)
                sc=f'{get_column_letter(col)}{row["excel_row"]}';ec=f'{get_column_letter(col+1)}{row["excel_row"]}'
                lines.append(f'| {row["excel_row"]} | {row["Age"]} | CE{v["number"]} | `{xml_cells[sc].find("s:v",ns).text}` | `{xml_cells[ec].find("s:v",ns).text}` | {clock(v["start_day"])} → {clock(v["end_day"])} | {v["duration_s"]:g} |')
    lines += ['', '### Interval longer than its entire acquisition', '',
              '`210406A_F`, Excel row 7, CE1: `F7 = '
              + xml_cells['F7'].find('s:v',ns).text + '`, `G7 = '
              + xml_cells['G7'].find('s:v',ns).text + '`. '
              'Converted times: **12:15:50 → 15:15:58**; duration **10 808 s**; '
              'recording duration **1454 s**. This cannot fit even after subtracting a constant '
              'clock offset. No correction has been applied.', '',
              'Additional issues: `200625A_F` has missing count and no intervals (not confirmed '
              'seizure-free); isolated totals 90 at D25 and 180 at D1048574; blank ID header; '
              'spaces around `200625A-F`; inconsistent beginning-column spellings '
              '`DebutCE5` and `DebtCE11`; physical units and absolute acquisition start unavailable.', '',
              '## 5. Channels per recording and intersection', '',
              'Exact channel names and original order:', '']
    for rid, rec in recordings.items():
        lines.append(f'- **{rid} ({len(rec["channels"])} channels):** ' + ', '.join(f'`{c}`' for c in rec['channels']))
    common=sorted(set.intersection(*(set(r['channels']) for r in recordings.values())))
    lines += ['', f'**Common to all 22 recordings ({len(common)} channels):** '+', '.join(f'`{c}`' for c in common)+'.', '',
              '**EEG-only intersection (8 channels):** '+', '.join(f'`{c}`' for c in common if c.startswith('EEG'))+'.', '',
              'The common set includes auxiliary `ECG` and `SLI`; EMG names differ or are absent. '
              'This list is an inventory, not a channel-selection decision.', '',
              '## Conclusion and checkpoint', '',
              'The direct elapsed-time hypothesis fails for every annotated acquisition. '
              'It remains plausible that the numbers are clock times, but this test does not '
              'establish the missing acquisition offset. The next reliable input would be '
              'original recording start metadata or annotation-author clarification. '
              'No alignment is accepted, no annotation is fixed, and no dataset has been built. '
              'Awaiting user review before further work.', '']
    (DOCS / 'TIME_ALIGNMENT.md').write_text('\n'.join(lines))
    (DOCS / 'alignment_elapsed_test.json').write_text(json.dumps(tests,indent=2)+'\n')
    print(f'Consistent: {yes}; inconsistent: {no}; untestable: 1')
    print('Report:',DOCS/'TIME_ALIGNMENT.md')
    print('Figures:',FIGS)


if __name__ == '__main__':
    main()
